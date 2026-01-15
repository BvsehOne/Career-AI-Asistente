import streamlit as st
import google.generativeai as genai
import os
from dotenv import load_dotenv
import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document
from fpdf import FPDF
from gtts import gTTS
import io
import re
import sqlite3
import hashlib
import time

# --- 1. CONFIGURACIÓN ---
load_dotenv()
st.set_page_config(page_title="ScanMatch | ATS Resume Checker", layout="wide", page_icon="🎯")

# --- 2. GESTIÓN DE ESTADO ---
if 'analyzed' not in st.session_state: st.session_state['analyzed'] = False
if 'analysis_result' not in st.session_state: st.session_state['analysis_result'] = None
if 'pdf_data' not in st.session_state: st.session_state['pdf_data'] = None
if 'audio_data' not in st.session_state: st.session_state['audio_data'] = None
if 'score' not in st.session_state: st.session_state['score'] = 0
if 'cv_content' not in st.session_state: st.session_state['cv_content'] = ""
if 'job_content' not in st.session_state: st.session_state['job_content'] = ""
if 'logged_in' not in st.session_state: st.session_state['logged_in'] = False
if 'username' not in st.session_state: st.session_state['username'] = ""
if 'nombre_usuario' not in st.session_state: st.session_state['nombre_usuario'] = ""

# --- 3. BASE DE DATOS ---
def init_db():
    conn = sqlite3.connect('usuarios.db')
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users (username TEXT PRIMARY KEY, password TEXT, name TEXT, email TEXT)''')
    conn.commit()
    conn.close()

def crear_usuario(username, password, name, email):
    conn = sqlite3.connect('usuarios.db')
    c = conn.cursor()
    pwd_hash = hashlib.sha256(password.encode()).hexdigest()
    try:
        c.execute('INSERT INTO users VALUES (?,?,?,?)', (username, pwd_hash, name, email))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()

def verificar_login(username, password):
    conn = sqlite3.connect('usuarios.db')
    c = conn.cursor()
    pwd_hash = hashlib.sha256(password.encode()).hexdigest()
    c.execute('SELECT * FROM users WHERE username = ? AND password = ?', (username, pwd_hash))
    user = c.fetchone()
    conn.close()
    return user

init_db()

# --- 4. API & SIDEBAR ---
with st.sidebar:
    st.header("⚙️ Configuración")
    try:
        api_key = st.secrets["GOOGLE_API_KEY"]
    except:
        api_key = os.getenv("GOOGLE_API_KEY")
    
    if not api_key:
        st.error("❌ Falta API Key")
        st.stop()
    
    genai.configure(api_key=api_key)
    
    if st.session_state['logged_in']:
        st.success(f"Usuario: {st.session_state['nombre_usuario']}")
        if st.button("Cerrar Sesión"):
            st.session_state['logged_in'] = False
            st.rerun()
    else:
        st.caption("Inicia sesión para usar la herramienta.")

    st.divider()
    if st.button("🔄 Nueva Búsqueda (Reset)"):
        keys_to_keep = ['logged_in', 'username', 'nombre_usuario']
        for key in list(st.session_state.keys()):
            if key not in keys_to_keep:
                del st.session_state[key]
        st.rerun()

# --- 5. FUNCIONES CORE ---
def generar_contenido_seguro(prompt):
    modelos_a_probar = ["gemini-2.5-flash", "gemini-1.5-flash", "gemini-pro"]
    errores = []
    for nombre_modelo in modelos_a_probar:
        try:
            model = genai.GenerativeModel(nombre_modelo)
            response = model.generate_content(prompt)
            return response
        except Exception as e:
            if "429" in str(e): time.sleep(5)
            errores.append(f"{nombre_modelo}: {str(e)}")
            continue 
    raise Exception(f"Error AI: {errores}")

def limpiar_texto_audio(texto):
    """Elimina markdown para que el audio suene natural"""
    texto = re.sub(r'[\*#_`~]', '', texto) # Quita simbolos
    texto = re.sub(r'^\s*-\s+', '', texto, flags=re.MULTILINE) # Quita guiones de lista
    texto = texto.replace("Score:", "Puntaje:").replace("/", " de ")
    return texto

# --- 6. ESTILOS CSS ---
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Open+Sans:wght@400;600;700;800&display=swap');
    
    /* ANIMACIÓN DE FONDO */
    @keyframes gradient-animation {
        0% { background-position: 0% 50%; }
        50% { background-position: 100% 50%; }
        100% { background-position: 0% 50%; }
    }

    .stApp {
        background: linear-gradient(-45deg, #004F9F, #2563EB, #00d2ff, #F8C43A);
        background-size: 400% 400%;
        animation: gradient-animation 15s ease infinite;
    }
    
    /* CAPA DE VIDRIO (TRANSPARENTE) */
    .stApp::before {
        content: ""; position: absolute; top: 0; left: 0; width: 100%; height: 100%;
        background: rgba(255, 255, 255, 0.4); 
        backdrop-filter: blur(12px);
        z-index: -1;
    }

    html, body, [class*="css"] {
        font-family: 'Open Sans', sans-serif;
        color: #333;
    }

    /* NAVBAR */
    .navbar {
        background: rgba(255, 255, 255, 0.9);
        padding: 15px 40px; border-bottom: 1px solid #ddd;
        display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; border-radius: 8px;
        backdrop-filter: blur(5px);
    }
    .brand { font-size: 24px; font-weight: 800; color: #004F9F; }
    
    /* CARDS */
    .scanner-card {
        background: rgba(255, 255, 255, 0.95);
        border-radius: 8px; box-shadow: 0 8px 32px 0 rgba(31, 38, 135, 0.15);
        padding: 30px; border: 1px solid rgba(255, 255, 255, 0.18); margin-bottom: 30px;
        backdrop-filter: blur(4px);
    }
    .login-box {
        max-width: 400px; margin: 50px auto; padding: 40px; 
        background: rgba(255, 255, 255, 0.95);
        border-radius: 12px; box-shadow: 0 8px 32px 0 rgba(31, 38, 135, 0.2); text-align: center;
        border: 1px solid rgba(255, 255, 255, 0.18);
    }
    
    /* BOTONES */
    .stButton > button {
        background-color: #F8C43A; color: #111; font-weight: 700; border: none;
        padding: 12px 30px; font-size: 18px; width: 100%; text-transform: uppercase;
        border-radius: 4px;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
    }
    .stButton > button:hover { background-color: #E5B020; transform: translateY(-1px); box-shadow: 0 6px 8px rgba(0,0,0,0.15); }
    
    /* SCORE FALLBACK (CSS) */
    .score-circle {
        width: 180px; height: 180px; border-radius: 50%;
        display: flex; flex-direction: column; justify-content: center; align-items: center;
        margin: 0 auto; border: 10px solid #ddd; background: white;
    }
    .icon-box { font-size: 30px; margin-bottom: 10px; }
</style>
""", unsafe_allow_html=True)

# --- 7. UTILIDADES PDF/WEB ---
class PDFJobScan(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 20)
        self.set_text_color(0, 79, 159)
        self.cell(0, 10, 'SCANMATCH REPORT', 0, 1, 'C')
        self.ln(5)
    def section_title(self, label):
        self.set_font('Arial', 'B', 12)
        self.set_fill_color(240, 240, 240)
        self.set_text_color(0, 0, 0)
        self.cell(0, 8, label, 0, 1, 'L', fill=True)
        self.ln(2)
    def body_text(self, text):
        self.set_font('Arial', '', 10)
        safe_text = text.encode('latin-1', 'replace').decode('latin-1')
        self.multi_cell(0, 5, safe_text)
        self.ln(2)

def generar_pdf(contenido):
    pdf = PDFJobScan()
    pdf.add_page()
    lines = contenido.split('\n')
    for line in lines:
        line = line.strip()
        if not line: continue
        if line.startswith("###"): pdf.section_title(line.replace("#","").strip())
        else: pdf.body_text(line)
    return pdf.output(dest='S').encode('latin-1')

def leer_web(url):
    try:
        r = requests.get(url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=5)
        soup = BeautifulSoup(r.text, 'html.parser')
        for t in soup(["script", "style", "nav", "footer"]): t.decompose()
        return soup.get_text(separator=' ', strip=True)
    except: return ""

def leer_doc(archivo):
    text = ""
    if archivo.name.endswith(".pdf"):
        reader = PdfReader(archivo)
        for p in reader.pages: text += p.extract_text()
    elif archivo.name.endswith(".docx"):
        doc = Document(archivo)
        text = "\n".join([p.text for p in doc.paragraphs])
    return text

# --- 8. LÓGICA DE FLUJO PRINCIPAL ---

# PANTALLA A: LOGIN
if not st.session_state['logged_in']:
    st.markdown("<h1 style='text-align:center; color:#004F9F;'>ScanMatch<span style='color:#F8C43A'>.io</span></h1>", unsafe_allow_html=True)
    
    col_l1, col_l2, col_l3 = st.columns([1,1,1])
    with col_l2:
        st.markdown('<div class="login-box">', unsafe_allow_html=True)
        tabs = st.tabs(["Ingresar", "Registrarse"])
        
        with tabs[0]:
            u = st.text_input("Usuario", key="u_in")
            p = st.text_input("Contraseña", type="password", key="p_in")
            if st.button("Iniciar Sesión"):
                user = verificar_login(u, p)
                if user:
                    st.session_state['logged_in'] = True
                    st.session_state['username'] = u
                    st.session_state['nombre_usuario'] = user[2]
                    st.rerun()
                else: st.error("Datos incorrectos")
        
        with tabs[1]:
            nu = st.text_input("Usuario Nuevo", key="u_new")
            np = st.text_input("Contraseña Nueva", type="password", key="p_new")
            nn = st.text_input("Nombre Completo", key="n_new")
            ne = st.text_input("Email", key="e_new")
            if st.button("Crear Cuenta"):
                if crear_usuario(nu, np, nn, ne): st.success("¡Creado! Inicia sesión.")
                else: st.error("Usuario ya existe")
        st.markdown('</div>', unsafe_allow_html=True)

# PANTALLA B: APP PRINCIPAL
else:
    st.markdown("""
    <div class="navbar">
        <div class="brand">ScanMatch<span style="color:#F8C43A">.io</span></div>
        <div style="font-weight:600; color:#004F9F;">Perfil Verificado</div>
    </div>
    """, unsafe_allow_html=True)

    # VISTA 1: INPUTS
    if not st.session_state['analyzed']:
        st.markdown("<h1 style='text-align:center;'>Escanee su currículum hoy</h1>", unsafe_allow_html=True)
        st.markdown("<p style='text-align:center;'>Descubra cómo hacerlo destacar ante los empleadores.</p>", unsafe_allow_html=True)

        with st.container():
            st.markdown('<div class="scanner-card">', unsafe_allow_html=True)
            col1, col2 = st.columns(2)
            
            with col1:
                st.markdown("### 1. Sube tu currículum")
                f_cv = st.file_uploader("Sube tu CV (PDF/DOCX)", label_visibility="collapsed")
                t_cv = ""
                if f_cv:
                    t_cv = leer_doc(f_cv)
                    st.success("✅ Cargado")

            with col2:
                st.markdown("### 2. Pegue la oferta de trabajo")
                tabs = st.tabs(["Texto", "URL"])
                t_job = ""
                with tabs[0]:
                    txt = st.text_area("Descripción", height=150, label_visibility="collapsed")
                    if txt: t_job = txt
                with tabs[1]:
                    url = st.text_input("Link Oferta", label_visibility="collapsed")
                    if url: t_job = leer_web(url)

            st.markdown("<br>", unsafe_allow_html=True)
            c1, c2, c3 = st.columns([1, 2, 1])
            with c2:
                if st.button("ESCANEA MI CURRÍCULUM"):
                    if t_cv and t_job:
                        with st.spinner("Analizando..."):
                            st.session_state['cv_content'] = t_cv
                            st.session_state['job_content'] = t_job
                            
                            prompt = f"""
                            Analiza este CV contra esta Oferta como un experto ATS.
                            Usa Markdown.
                            1. SCORE: (0-100)
                            2. RESUMEN: Diagnóstico breve.
                            3. HABILIDADES DURAS: Lista las presentes y faltantes.
                            4. HABILIDADES BLANDAS: Lista las presentes y faltantes.
                            5. CHEQUEO ATS: Evalúa formato, fecha, imágenes.
                            6. CONSEJOS: 3 tips de reclutador.

                            CV: {t_cv[:15000]}
                            OFERTA: {t_job[:15000]}
                            """
                            try:
                                res = generar_contenido_seguro(prompt)
                                st.session_state['analysis_result'] = res.text
                                try:
                                    score_match = re.search(r"SCORE:\s*(\d+)", res.text)
                                    st.session_state['score'] = int(score_match.group(1)) if score_match else 50
                                except: st.session_state['score'] = 50
                                
                                st.session_state['analyzed'] = True
                                st.rerun()
                            except Exception as e:
                                st.error(f"Error AI: {e}")
                    else:
                        st.warning("Sube ambos archivos.")
            st.markdown('</div>', unsafe_allow_html=True)

        # MARKETING
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown("<h2 style='text-align:center;'>Nuestro verificador de currículums busca:</h2>", unsafe_allow_html=True)
        
        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.markdown("""<div class="scanner-card"><div class="icon-box">🔍</div><h3>Consejos específicos de ATS</h3><p>Detecta el ATS específico para ajustar rendimiento.</p></div>""", unsafe_allow_html=True)
        with col_m2:
            st.markdown("""<div class="scanner-card"><div class="icon-box">📝</div><h3>Estilo y formato</h3><p>Tipo de archivo, consistencia y legibilidad.</p></div>""", unsafe_allow_html=True)
        col_m3, col_m4 = st.columns(2)
        with col_m3:
            st.markdown("""<div class="scanner-card"><div class="icon-box">📊</div><h3>Palabras clave y habilidades</h3><p>Habilidades duras, blandas y faltantes críticas.</p></div>""", unsafe_allow_html=True)
        with col_m4:
            st.markdown("""<div class="scanner-card"><div class="icon-box">✨</div><h3>Contenido del CV</h3><p>Coincidencia de títulos, experiencia y logros.</p></div>""", unsafe_allow_html=True)
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown("<h2 style='text-align:center;'>Cómo utilizar el escáner</h2>", unsafe_allow_html=True)
        col_s1, col_s2, col_s3, col_s4 = st.columns(4)
        with col_s1: st.markdown("<h1 style='color:#F8C43A; text-align:center;'>1</h1><p style='text-align:center; font-weight:bold;'>Sube tu currículum</p>", unsafe_allow_html=True)
        with col_s2: st.markdown("<h1 style='color:#F8C43A; text-align:center;'>2</h1><p style='text-align:center; font-weight:bold;'>Pega la oferta</p>", unsafe_allow_html=True)
        with col_s3: st.markdown("<h1 style='color:#F8C43A; text-align:center;'>3</h1><p style='text-align:center; font-weight:bold;'>Optimiza</p>", unsafe_allow_html=True)
        with col_s4: st.markdown("<h1 style='color:#F8C43A; text-align:center;'>4</h1><p style='text-align:center; font-weight:bold;'>Aumenta Match</p>", unsafe_allow_html=True)

    # VISTA 2: RESULTADOS
    else:
        score = st.session_state['score']
        
        st.markdown(f"""
        <div style="background:white; padding:20px; text-align:center; border-bottom:1px solid #ddd; margin-bottom:30px; border-radius:8px;">
            <h2>Resultados del Análisis ATS</h2>
        </div>
        """, unsafe_allow_html=True)
        
        # --- LÓGICA DE GRÁFICO (CON O SIN PLOTLY) ---
        try:
            import plotly.graph_objects as go
            fig = go.Figure(go.Indicator(
                mode = "gauge+number",
                value = score,
                domain = {'x': [0, 1], 'y': [0, 1]},
                title = {'text': "Match Rate"},
                gauge = {
                    'axis': {'range': [0, 100], 'tickwidth': 1, 'tickcolor': "darkblue"},
                    'bar': {'color': "#004F9F"},
                    'bgcolor': "white",
                    'borderwidth': 2,
                    'bordercolor': "gray",
                    'steps': [
                        {'range': [0, 50], 'color': '#FFCCCC'},
                        {'range': [50, 75], 'color': '#FFF4CC'},
                        {'range': [75, 100], 'color': '#CCFFCC'}],
                    'threshold': {
                        'line': {'color': "red", 'width': 4},
                        'thickness': 0.75,
                        'value': score}}))
            fig.update_layout(height=300, margin=dict(l=20, r=20, t=30, b=20), paper_bgcolor="rgba(0,0,0,0)", font={'color': "black"})
            st.plotly_chart(fig, use_container_width=True)
        except ImportError:
            # FALLBACK SI NO INSTALÓ PLOTLY
            color = "#2ECC71" if score > 75 else ("#F1C40F" if score > 50 else "#E74C3C")
            st.markdown(f"""
            <div class="score-circle" style="border-color:{color};">
                <span style="font-size:3.5rem; font-weight:800; color:#333;">{score}%</span>
                <span style="font-size:1rem; color:#666; font-weight:600;">MATCH RATE</span>
            </div>
            <p style="text-align:center; color:#666;">(Instala 'plotly' para ver el gráfico Pro)</p>
            """, unsafe_allow_html=True)

        with st.container():
            st.markdown('<div class="scanner-card">', unsafe_allow_html=True)
            t1, t2, t3, t4 = st.tabs(["🔍 Palabras Clave", "⚙️ Formato", "💡 Consejos", "📥 Herramientas"])
            
            raw = st.session_state['analysis_result']
            
            with t1:
                st.info("Palabras clave encontradas vs faltantes:")
                try:
                    part = raw.split("HABILIDADES DURAS")[1].split("CHEQUEO ATS")[0]
                    clean_part = part.replace("4.", "").replace("5.", "").replace("6.", "")
                    st.markdown("### Habilidades Duras y Blandas")
                    st.markdown(clean_part)
                except: st.write(raw)
                
            with t2:
                st.info("Análisis de formato ATS:")
                try:
                    part = raw.split("CHEQUEO ATS")[1].split("CONSEJOS")[0]
                    st.markdown(part)
                except: st.write("Ver reporte completo.")
                
            with t3:
                st.success("Recomendaciones de experto:")
                try:
                    part = raw.split("CONSEJOS")[1]
                    st.markdown(part)
                except: st.write("Ver reporte completo.")
                
            with t4:
                c_d1, c_d2 = st.columns(2)
                with c_d1:
                    if st.button("📄 Generar Reporte PDF"):
                        st.session_state['pdf_data'] = generar_pdf(raw)
                    if st.session_state['pdf_data']:
                        st.download_button("Descargar PDF", st.session_state['pdf_data'], "Reporte.pdf", "application/pdf")
                with c_d2:
                    if st.button("🎧 Generar Audio Entrevista"):
                        if st.session_state['job_content']:
                            try:
                                with st.spinner("Creando entrevista..."):
                                    prompt_audio = f"Actúa como reclutador para este puesto: {st.session_state['job_content'][:2000]}. Hazme una pregunta difícil y breve. NO USES ASTERISCOS NI FORMATO."
                                    q = generar_contenido_seguro(prompt_audio)
                                    
                                    # LIMPIEZA DE AUDIO MEJORADA
                                    texto_limpio = limpiar_texto_audio(q.text)
                                    
                                    tts = gTTS(texto_limpio, lang='es')
                                    buf = io.BytesIO()
                                    tts.write_to_fp(buf)
                                    buf.seek(0)
                                    st.session_state['audio_data'] = buf
                                    st.write(f"🗣️ **Pregunta:** {texto_limpio}")
                            except Exception as e: st.error(f"Error audio: {e}")
                        else: st.error("Error de contexto.")
                        
                    if st.session_state['audio_data']:
                        st.audio(st.session_state['audio_data'], format='audio/mp3')

            st.markdown('</div>', unsafe_allow_html=True)